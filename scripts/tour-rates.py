"""ADR-097: read the filled-in tour rate sheet (docs/imports/tour-rates-template.xlsx) as JSON for pnpm import:tour-rates.

    /usr/local/bin/python3 -I scripts/tour-rates.py path/to/tour-rates.xlsx > rates.json

Prints each tab's rows as they are (example rows and rows with no prices left out); scripts/lib/tour-rates.ts checks
and converts them. Formula columns (the grey totals) are not read.
"""
import datetime
import json
import sys

import openpyxl

TABS = {
    "Shared tours": ["tour", "office", "days", "times", "adult", "child611", "child35", "max", "min", "lead", "tickets", "notes"],
    "Private tours": ["tour", "office", "days", "times", "p1", "p2", "p34", "p58", "p914", "p15", "child_pct", "lead", "notes"],
    "Transfers": ["tour", "office", "label", "car", "van", "minibus", "night", "lead", "notes"],
    "Meet and assist": ["tour", "office", "label", "t14", "t58", "t914", "lead", "notes"],
    "Seasons": ["name", "from", "to", "pct", "applies", "notes"],
    "Questions": ["n", "question", "answer", "recommended"],
}
HEADERS = {"Tour", "Transfer", "Service", "Season", "#"}
PRICES = {"adult", "p1", "p2", "p34", "p58", "p914", "p15", "car", "van", "minibus", "t14", "t58", "t914"}


def cell(v):
    if isinstance(v, datetime.datetime):
        return v.date().isoformat() if (v.hour, v.minute) == (0, 0) else v.strftime("%H:%M")
    if isinstance(v, (datetime.date,)):
        return v.isoformat()
    if isinstance(v, datetime.time):
        return v.strftime("%H:%M")
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


wb = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)
out = {}
for tab, cols in TABS.items():
    rows = []
    if tab in wb.sheetnames:
        seen_header = False
        for n, r in enumerate(wb[tab].iter_rows(values_only=True), start=1):
            row = {c: cell(v) for c, v in zip(cols, r)}
            first = row.get(cols[0])
            if not seen_header:  # rows start under the header ("Tour", "Transfer", "Service", "Season", "#")
                seen_header = first in HEADERS
                continue
            if first is None or str(first).lower().startswith("example"):
                continue
            if tab not in ("Seasons", "Questions") and not any(row.get(k) is not None for k in PRICES):
                continue
            if tab == "Questions" and row.get("answer") is None:
                continue
            rows.append({"row": n, **row})
    out[tab] = rows
json.dump(out, sys.stdout, ensure_ascii=False, indent=1)
print()
